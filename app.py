from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

import fitz  # PyMuPDF
import openpyxl
from flask import Flask, jsonify, render_template, request, send_file
from pdf2docx import Converter
from PIL import Image

app = Flask(__name__)

SUPPORTED_CONVERSIONS = {
    "doc_to_pdf": {"input": ["doc", "docx"], "output": "pdf"},
    "excel_to_pdf": {"input": ["xls", "xlsx"], "output": "pdf"},
    "image_to_pdf": {"input": ["jpg", "jpeg", "png"], "output": "pdf"},
    "pdf_to_image": {"input": ["pdf"], "output": ["jpg", "jpeg", "png"]},
    "pdf_to_docx": {"input": ["pdf"], "output": "docx"},
    "pdf_to_xlsx": {"input": ["pdf"], "output": "xlsx"},
}


def _suffix(file_name: str) -> str:
    return Path(file_name).suffix.lower().lstrip(".")


def _save_upload(upload, target_dir: Path) -> Path:
    safe_name = f"{uuid.uuid4().hex}_{Path(upload.filename).name}"
    file_path = target_dir / safe_name
    upload.save(file_path)
    return file_path


def _convert_office_to_pdf(input_path: Path, output_dir: Path) -> Path:
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(input_path),
        ],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    pdf_file = output_dir / f"{input_path.stem}.pdf"
    if not pdf_file.exists():
        raise RuntimeError("Office to PDF conversion failed.")
    return pdf_file


def _convert_image_to_pdf(input_path: Path, output_dir: Path) -> Path:
    image = Image.open(input_path)
    rgb = image.convert("RGB")
    output_path = output_dir / f"{input_path.stem}.pdf"
    rgb.save(output_path, "PDF", resolution=300.0)
    return output_path


def _convert_pdf_to_images(input_path: Path, output_dir: Path, image_format: str) -> Path:
    pdf_doc = fitz.open(input_path)
    generated_files = []
    for page_no, page in enumerate(pdf_doc, start=1):
        pix = page.get_pixmap(dpi=200)
        out = output_dir / f"{input_path.stem}_page_{page_no}.{image_format}"
        pix.save(out)
        generated_files.append(out)

    if len(generated_files) == 1:
        return generated_files[0]

    zip_path = output_dir / f"{input_path.stem}_images.zip"
    shutil.make_archive(str(zip_path.with_suffix("")), "zip", output_dir)
    return zip_path


def _convert_pdf_to_docx(input_path: Path, output_dir: Path) -> Path:
    output_path = output_dir / f"{input_path.stem}.docx"
    converter = Converter(str(input_path))
    converter.convert(str(output_path), start=0, end=None)
    converter.close()
    return output_path


def _convert_pdf_to_xlsx(input_path: Path, output_dir: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Extracted Text"

    pdf_doc = fitz.open(input_path)
    row = 1
    for page_index, page in enumerate(pdf_doc, start=1):
        sheet.cell(row=row, column=1, value=f"Page {page_index}")
        row += 1
        lines = page.get_text("text").splitlines()
        for line in lines:
            sheet.cell(row=row, column=1, value=line)
            row += 1
        row += 1

    output_path = output_dir / f"{input_path.stem}.xlsx"
    workbook.save(output_path)
    return output_path


def _redact_pdf(input_path: Path, output_dir: Path, terms: list[str]) -> Path:
    doc = fitz.open(input_path)
    for page in doc:
        for term in terms:
            rects = page.search_for(term)
            for rect in rects:
                page.add_redact_annot(rect, fill=(0, 0, 0))
        page.apply_redactions()

    output_path = output_dir / f"{input_path.stem}_redacted.pdf"
    doc.save(output_path)
    doc.close()
    return output_path


@app.route("/")
def index():
    return render_template("index.html", conversions=SUPPORTED_CONVERSIONS)


@app.route("/api/convert", methods=["POST"])
def convert():
    if "file" not in request.files:
        return jsonify({"error": "Missing file upload."}), 400

    mode = request.form.get("mode", "")
    target_format = request.form.get("target_format", "pdf").lower()
    uploaded = request.files["file"]

    if mode not in SUPPORTED_CONVERSIONS:
        return jsonify({"error": "Unsupported conversion mode."}), 400

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        source = _save_upload(uploaded, tmp_dir)

        input_ext = _suffix(uploaded.filename)
        expected_inputs = SUPPORTED_CONVERSIONS[mode]["input"]
        if input_ext not in expected_inputs:
            return jsonify({"error": f"Expected one of: {expected_inputs}"}), 400

        if mode == "doc_to_pdf" or mode == "excel_to_pdf":
            output_path = _convert_office_to_pdf(source, tmp_dir)
        elif mode == "image_to_pdf":
            output_path = _convert_image_to_pdf(source, tmp_dir)
        elif mode == "pdf_to_image":
            if target_format not in ["jpg", "jpeg", "png"]:
                return jsonify({"error": "PDF to image supports jpg/jpeg/png."}), 400
            output_path = _convert_pdf_to_images(source, tmp_dir, target_format)
        elif mode == "pdf_to_docx":
            output_path = _convert_pdf_to_docx(source, tmp_dir)
        elif mode == "pdf_to_xlsx":
            output_path = _convert_pdf_to_xlsx(source, tmp_dir)
        else:
            return jsonify({"error": "Unsupported conversion mode."}), 400

        output_bytes = output_path.read_bytes()
        output_name = output_path.name

    return send_file(
        io.BytesIO(output_bytes),
        as_attachment=True,
        download_name=output_name,
        mimetype="application/octet-stream",
    )


@app.route("/api/redact", methods=["POST"])
def redact():
    if "file" not in request.files:
        return jsonify({"error": "Missing file upload."}), 400

    uploaded = request.files["file"]
    if _suffix(uploaded.filename) != "pdf":
        return jsonify({"error": "Redaction only supports PDF files."}), 400

    terms_raw = request.form.get("terms", "")
    terms = [x.strip() for x in terms_raw.split(",") if x.strip()]
    if not terms:
        return jsonify({"error": "Provide comma-separated terms to redact."}), 400

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        source = _save_upload(uploaded, tmp_dir)
        output_path = _redact_pdf(source, tmp_dir, terms)
        output_bytes = output_path.read_bytes()
        output_name = output_path.name

    return send_file(
        io.BytesIO(output_bytes),
        as_attachment=True,
        download_name=output_name,
        mimetype="application/pdf",
    )


if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=True)
